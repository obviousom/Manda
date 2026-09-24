#!/usr/bin/env python3
"""Append one row to daily_summary.xlsx from a `claude -p --output-format
json --json-schema ...` result. Called by summarize-daily-log.sh.

Usage: append-summary-row.py <xlsx_path> <date_str> <result_json>
Exits non-zero (and prints why) without touching the file if the result
doesn't parse - caller keeps the pending MD around for the next retry.
"""
import sys
import json
import openpyxl


def main():
    if len(sys.argv) != 4:
        print("usage: append-summary-row.py <xlsx_path> <date_str> <result_json>")
        sys.exit(2)

    xlsx_path, date_str, result_json = sys.argv[1], sys.argv[2], sys.argv[3]

    try:
        data = json.loads(result_json)
    except Exception as e:
        print(f"could not parse claude result as JSON: {e}")
        sys.exit(1)

    so = data.get("structured_output")
    if not so:
        print(f"no structured_output in claude result: {result_json[:500]}")
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    next_row = ws.max_row + 1

    ws.cell(row=next_row, column=1, value=date_str)
    ws.cell(row=next_row, column=2, value=so.get("summary") or "—")
    ws.cell(row=next_row, column=3, value=so.get("key_outcomes") or "—")
    ws.cell(row=next_row, column=4, value=so.get("issues_filed") or "—")
    ws.cell(row=next_row, column=5, value=so.get("branches_prs") or "—")
    ws.cell(row=next_row, column=6, value=so.get("notes") or "—")

    wb.save(xlsx_path)
    print(f"appended row {next_row} for {date_str}")


if __name__ == "__main__":
    main()
